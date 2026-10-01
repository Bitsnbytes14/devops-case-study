"""Generate diagrams, case study report, and five slide presentation."""
from pathlib import Path
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).parents[1]; DOCS = ROOT / 'docs'; DIAGRAMS = DOCS / 'diagrams'
DIAGRAMS.mkdir(parents=True, exist_ok=True)

def diagram(name, labels, color):
    fig, ax = plt.subplots(figsize=(11, 3)); ax.axis('off'); ax.set_xlim(0, len(labels)); ax.set_ylim(0, 1)
    for i, label in enumerate(labels):
        ax.text(i + .5, .5, label, ha='center', va='center', fontsize=11, bbox=dict(boxstyle='round,pad=.6', facecolor=color, edgecolor='#17324d'))
        if i < len(labels) - 1: ax.annotate('', (i + .95, .5), (i + 1.05, .5), arrowprops={'arrowstyle': '->', 'lw': 2})
    fig.tight_layout(); fig.savefig(DIAGRAMS / name, dpi=180); plt.close(fig)

diagram('pipeline.png', ['Test', 'Security', 'Build and push', 'Kind deploy', 'Smoke test'], '#9dd9d2')
diagram('architecture.png', ['Students', 'Flask API', 'Allocation engine', 'Prometheus', 'Grafana'], '#b8d8f8')
diagram('netflix-case.png', ['Monolith', 'AWS services', 'Circuit breakers', 'Chaos experiments'], '#ffc9a5')
diagram('amazon-case.png', ['Obidos', 'Service teams', 'Automated delivery', 'Fast feedback'], '#d5c6e8')

netflix = '''Netflix faced a defining reliability event in 2008 when database corruption affected its DVD operation. The incident exposed the limits of a monolithic application connected to a vertically scaled central database. Recovery and change were difficult because one failure domain carried customer experience, catalog operations, and operational risk. Scaling the same machine could improve capacity, but it could not isolate faults or make independent releases safer. The company needed a system that could grow in demand while remaining available during component failure.

Netflix moved to AWS and decomposed the monolith into microservices. Each service could own its data and be deployed independently, reducing the blast radius of a failure. Database per service boundaries limited coupling. Circuit breakers allowed callers to degrade gracefully when a dependency failed instead of exhausting resources through repeated retries. Spinnaker brought repeatable deployment practices, with automated promotion and rollback controls. The operational model embraced resilience testing through Chaos Monkey, the Simian Army, and larger experiments such as Chaos Kong. These tools deliberately introduced failures so teams could discover weak assumptions before customers encountered them.

The outcome was a platform designed for failure rather than dependent on avoiding every failure. Teams gained the ability to release services independently, scale selected workloads, and restore service quickly. The key lesson is that availability is a product property supported by architecture, automation, observability, and culture. Resilience comes from many small protections that are continuously tested.

RoomFit demonstrates the same ideas at a smaller scale. Its allocation API is independently deployable, replicated three times, and protected by readiness and liveness probes. Rolling updates retain availability while a new image is introduced. Prometheus measures latency and error behavior, while the fault rate and added latency switches make controlled chaos testing possible. These choices do not recreate Netflix scale, but they apply the same principle: verify that an important service behaves predictably when change and failure occur.'''
amazon = '''Amazon began with the Obidos monolith, an application whose closely connected components made releases risky and outages expensive. As the company expanded, teams had to coordinate changes through a shared codebase and shared operational constraints. This slowed delivery and made it difficult to identify ownership when a service degraded. A large release could contain many unrelated changes, so rollback carried uncertainty and operational work increased with every dependency.

Amazon responded by evolving toward a service oriented architecture. Small teams, often described as two pizza teams, owned focused services and their customer outcomes. The culture of you build it you run it joined development accountability with production responsibility. Teams could choose practical implementation details while retaining explicit contracts with other services. Automated deployment pipelines created a repeatable path from tested change to production, and rollback reduced the cost of a failed release. Amazon has reported deployment rates averaging one every 11.7 seconds, illustrating how decoupled ownership and automation can turn delivery into a routine capability.

The outcome was more than technical modularity. It was a culture of continuous innovation where teams could experiment, observe results, and improve quickly. Independent services reduced coordination cost, while shared platform practices gave teams consistent safety controls. The case shows that speed is sustainable only when it includes quality gates, monitoring, and clear ownership.

RoomFit applies this lesson through a narrowly scoped allocation microservice. The engine has one responsibility: apply compatibility scores while enforcing non negotiable room constraints. GitHub Actions validates code, security, image creation, deployment, and smoke tests. The service can be changed without a database dependency, and Kubernetes supplies replicas, probes, and rollback. The result is a practical demonstration of how independently deployable services, automated checks, and observable outcomes allow a team to make smaller, safer changes.'''

def report():
    d = Document(); d.add_heading('RoomFit DevOps Case Studies', 0); d.add_paragraph('CA II evaluation evidence')
    for title, text, image in [('Q1 Netflix reliability transformation', netflix, 'netflix-case.png'), ('Q2 Amazon service ownership transformation', amazon, 'amazon-case.png')]:
        d.add_heading(title, 1); d.add_heading('Problem, solution, outcome, and RoomFit connection', 2)
        for para in text.split('\n\n'): d.add_paragraph(para)
        d.add_picture(str(DIAGRAMS / image), width=Inches(6.2))
    d.add_heading('Comparison', 1); table = d.add_table(rows=1, cols=3); table.style = 'Light Shading Accent 1'
    for cell, text in zip(table.rows[0].cells, ['Case', 'Core practice', 'RoomFit evidence']): cell.text = text
    for row in [('Netflix', 'Resilience and chaos testing', 'Probes, replicas, chaos switches'), ('Amazon', 'Independent ownership and delivery', 'Focused API and automated pipeline')]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row): cell.text = value
    d.save(DOCS / 'RoomFit_DevOps_Case_Studies.docx')

def slides():
    p = Presentation(); p.slide_width = PInches(13.333); p.slide_height = PInches(7.5)
    slides = [('Architecture', 'Students call the Flask API. The deterministic engine enforces hard room constraints. Prometheus collects health and business metrics. Grafana presents the operational view.', 'architecture.png'), ('Pipeline flow', 'Test and coverage gates run first. Security scanning follows. Main branch changes build an image, deploy to Kind, and run smoke tests.', 'pipeline.png'), ('Challenges', 'Hard constraints must never be violated. Partial groups must still receive a room. Production safety requires probes, least privilege, resource limits, and observable failure behavior.', 'netflix-case.png'), ('Lessons learned', 'Small deployable services lower release risk. Automation makes safety repeatable. Monitoring reveals impact quickly. Controlled faults test assumptions before users find them.', 'amazon-case.png'), ('Results and evidence', 'Pytest proves allocation rules. Coverage gate targets 85 percent. Kubernetes manifests support rolling update and rollback. Grafana includes uptime, latency, error rate, version, and allocation metrics.', 'architecture.png')]
    for index, (title, body, image) in enumerate(slides):
        s = p.slides.add_slide(p.slide_layouts[6]); bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(18, 35, 54)
        box = s.shapes.add_textbox(PInches(.7), PInches(.55), PInches(6.2), PInches(.7)); tf = box.text_frame; tf.text = title; tf.paragraphs[0].font.size = Pt(36); tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = RGBColor(157, 217, 210)
        bodybox = s.shapes.add_textbox(PInches(.8), PInches(1.7), PInches(5.6), PInches(3.9)); btf = bodybox.text_frame; btf.word_wrap = True; btf.text = body; btf.paragraphs[0].font.size = Pt(22); btf.paragraphs[0].font.color.rgb = RGBColor(245, 248, 250)
        s.shapes.add_picture(str(DIAGRAMS / image), PInches(7), PInches(1.4), width=PInches(5.5))
        foot = s.shapes.add_textbox(PInches(.8), PInches(6.7), PInches(4), PInches(.3)); foot.text_frame.text = f'RoomFit DevOps case study  |  {index + 1} of 5'; foot.text_frame.paragraphs[0].font.color.rgb = RGBColor(180, 200, 210); foot.text_frame.paragraphs[0].font.size = Pt(11)
    p.save(DOCS / 'RoomFit_DevOps_Slides.pptx')

report(); slides()
